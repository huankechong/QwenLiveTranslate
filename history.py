"""翻译历史记录（SQLite，体验层功能包 2026-09-19）。

对标 LiveCaptionsTranslator 的翻译历史：
- 每句完整译文（final）落库：时间、源语种、目标语种、原文、译文、延迟
- 写入发生在 controller 回调线程，读发生在 UI 线程 → 用锁保护
- 提供 CSV 导出（UTF-8-BOM，Excel 直接打开不乱码）
- 库文件 translation_history.db 放 exe/脚本同目录（便携）

表结构：
  history(id INTEGER PRIMARY KEY AUTOINCREMENT,
          ts REAL,            -- unix 时间戳
          source_lang TEXT,   -- 源语种（auto=自动识别）
          target_lang TEXT,   -- 目标语种
          speaker TEXT,       -- 说话人标签（可空）
          original TEXT,      -- 原文（ASR 终稿）
          translation TEXT,   -- 译文（response.text.done 终稿）
          latency_ms INT)     -- 本句延迟（可空）
"""

from __future__ import annotations

import csv
import sqlite3
import sys
import threading
import time
from pathlib import Path


def _db_path() -> Path:
    """历史库路径跟随 settings 的目录策略（含只读目录回退，M1 同修）。

    打包成 exe 后 __file__ 指向临时解压目录；Program Files 等
    只读场景由 settings._base_dir 统一回退 %APPDATA%。
    """
    import settings as _st
    return _st._base_dir() / "translation_history.db"


def db_dir() -> Path:
    """历史库所在目录（供导出默认路径等使用；公开 API，勿用 _db_path）。"""
    return _db_path().parent


class _NullStore:
    """降级用空实现：库不可用（如 exe 放只读目录）时保住应用启动，
    历史功能静默停用而不是整个程序崩溃。"""

    def add(self, *a, **k):  # noqa: ANN001, ANN002, ANN003, D102
        return -1

    def recent(self, limit=200, offset=0):  # noqa: ANN001, ANN201, D102
        return []

    def count(self):  # noqa: ANN201, D102
        return 0

    def clear(self):  # noqa: D102
        pass

    def close(self):  # noqa: D102
        pass

    def export_csv(self, out_path, limit=None):  # noqa: ANN001, ANN201, D102
        raise RuntimeError("历史库不可用（目录只读？），无法导出")


class HistoryStore:
    """线程安全的 SQLite 历史仓库（WAL 模式，读写都走同一连接+锁）。"""

    def __init__(self, db_path: Path | None = None):
        self.path = Path(db_path) if db_path else _db_path()
        self._lock = threading.Lock()
        self._conn = sqlite3.connect(str(self.path), check_same_thread=False)
        self._conn.execute("PRAGMA journal_mode=WAL")
        self._conn.execute(
            """CREATE TABLE IF NOT EXISTS history(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ts REAL,
                source_lang TEXT,
                target_lang TEXT,
                speaker TEXT,
                original TEXT,
                translation TEXT,
                latency_ms INT)"""
        )
        self._conn.commit()

    # ---------- 写 ----------
    def add(self, source_lang: str, target_lang: str, speaker: str | None,
            original: str, translation: str, latency_ms: int | None = None) -> int:
        with self._lock:
            cur = self._conn.execute(
                "INSERT INTO history(ts,source_lang,target_lang,speaker,original,translation,latency_ms)"
                " VALUES(?,?,?,?,?,?,?)",
                (time.time(), source_lang, target_lang, speaker, original, translation, latency_ms),
            )
            self._conn.commit()
            return cur.lastrowid

    # ---------- 读 ----------
    def recent(self, limit: int = 200, offset: int = 0) -> list[dict]:
        with self._lock:
            rows = self._conn.execute(
                "SELECT id,ts,source_lang,target_lang,speaker,original,translation,latency_ms"
                " FROM history ORDER BY id DESC LIMIT ? OFFSET ?",
                (limit, offset),
            ).fetchall()
        keys = ("id", "ts", "source_lang", "target_lang", "speaker",
                "original", "translation", "latency_ms")
        return [dict(zip(keys, r)) for r in rows]

    def count(self) -> int:
        with self._lock:
            return self._conn.execute("SELECT COUNT(*) FROM history").fetchone()[0]

    # ---------- 清理 ----------
    def clear(self):
        with self._lock:
            self._conn.execute("DELETE FROM history")
            self._conn.commit()

    def close(self):
        with self._lock:
            try:
                self._conn.close()
            except Exception:  # noqa: BLE001
                pass

    # ---------- 导出 ----------
    def export_csv(self, out_path: str | Path, limit: int | None = None) -> int:
        """导出全部/最近 limit 条到 CSV（UTF-8-BOM），统一按时间正序（旧→新）。"""
        if limit:
            rows = self.recent(limit=limit)
            rows.reverse()  # recent 返回新→旧，反转统一为旧→新
        else:
            with self._lock:
                rows_raw = self._conn.execute(
                    "SELECT id,ts,source_lang,target_lang,speaker,original,translation,latency_ms"
                    " FROM history ORDER BY id ASC"
                ).fetchall()
            keys = ("id", "ts", "source_lang", "target_lang", "speaker",
                    "original", "translation", "latency_ms")
            rows = [dict(zip(keys, r)) for r in rows_raw]
        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        n = 0
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["时间", "源语种", "目标语种", "原文", "译文", "延迟ms"])
            for r in rows:
                w.writerow([
                    time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(r["ts"])),
                    r["source_lang"], r["target_lang"],
                    r["original"] or "", r["translation"] or "",
                    "" if r["latency_ms"] is None else r["latency_ms"],
                ])
                n += 1
        return n

    def _export_rows(self, limit: int | None) -> list[dict]:
        """导出用数据（时间正序 旧→新），CSV/xlsx 共用。"""
        if limit:
            rows = self.recent(limit=limit)
            rows.reverse()
            return rows
        with self._lock:
            rows_raw = self._conn.execute(
                "SELECT id,ts,source_lang,target_lang,speaker,original,translation,latency_ms"
                " FROM history ORDER BY id ASC"
            ).fetchall()
        keys = ("id", "ts", "source_lang", "target_lang", "speaker",
                "original", "translation", "latency_ms")
        return [dict(zip(keys, r)) for r in rows_raw]

    def export_xlsx(self, out_path: str | Path, limit: int | None = None) -> int:
        """导出为 Excel（.xlsx）：时间列是真日期单元格（yyyy-mm-dd hh:mm:ss），
        列宽/表头样式预设，Excel 打开即正确显示（CSV 会被 Excel 按区域
        设置自作主张转换/吞列宽，xlsx 才是显示可控的正式交付格式）。"""
        from datetime import datetime
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        rows = self._export_rows(limit)
        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "翻译历史"
        headers = ["时间", "源语种", "目标语种", "原文", "译文", "延迟ms"]
        ws.append(headers)
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="1F2937")
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")
        date_style = "yyyy-mm-dd hh:mm:ss"
        for r in rows:
            ts = r["ts"]
            ws.append([
                datetime.fromtimestamp(ts) if ts is not None else None,
                r["source_lang"] or "",
                r["target_lang"] or "",
                r["original"] or "",
                r["translation"] or "",
                "" if r["latency_ms"] is None else r["latency_ms"],
            ])
            ws.cell(row=ws.max_row, column=1).number_format = date_style
        for i, w_ in enumerate((19, 8, 8, 45, 45, 9), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w_
        wb.save(out)
        return len(rows)


# 模块级单例（console 启动时创建）
_STORE: HistoryStore | None = None
_STORE_LOCK = threading.Lock()


def get_store() -> HistoryStore | _NullStore:
    global _STORE
    with _STORE_LOCK:
        if _STORE is None:
            try:
                _STORE = HistoryStore()
            except Exception:  # noqa: BLE001
                # 目录不可写等：降级为空实现，保住应用启动
                _STORE = _NullStore()
        return _STORE


def shutdown():
    global _STORE
    with _STORE_LOCK:
        if _STORE is not None:
            _STORE.close()
            _STORE = None
