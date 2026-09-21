# -*- coding: utf-8 -*-
"""history 存储：增删查/CSV 导出/降级单例。"""
import csv
import time

import history


class TestHistoryStore:
    def test_add_recent_roundtrip(self, tmp_path):
        s = history.HistoryStore(tmp_path / "h.db")
        sid = s.add("zh", "de", None, "你好", "Hallo", 123)
        assert sid >= 1
        rows = s.recent(limit=10)
        assert len(rows) == 1
        assert rows[0]["original"] == "你好"
        assert rows[0]["translation"] == "Hallo"
        assert rows[0]["latency_ms"] == 123
        s.close()

    def test_recent_order_desc(self, tmp_path):
        s = history.HistoryStore(tmp_path / "h.db")
        for i in range(5):
            s.add("zh", "en", None, f"o{i}", f"t{i}")
        rows = s.recent(limit=3)
        assert [r["original"] for r in rows] == ["o4", "o3", "o2"]
        s.close()

    def test_clear(self, tmp_path):
        s = history.HistoryStore(tmp_path / "h.db")
        s.add("zh", "en", None, "a", "b")
        s.clear()
        assert s.count() == 0
        s.close()

    def test_export_csv_bom_and_order(self, tmp_path):
        s = history.HistoryStore(tmp_path / "h.db")
        s.add("zh", "en", None, "第一句", "first")
        s.add("zh", "en", None, "第二句", "second")
        out = tmp_path / "out.csv"
        n = s.export_csv(str(out))
        assert n == 2
        raw = out.read_bytes()
        assert raw[:3] == b"\xef\xbb\xbf"  # UTF-8 BOM
        rows = list(csv.reader(raw.decode("utf-8-sig").splitlines()))
        assert rows[0][:3] == ["时间", "源语种", "目标语种"]
        assert rows[1][3] == "第一句"  # 旧→新正序
        assert rows[2][3] == "第二句"
        s.close()

    def test_null_store_fallback(self, tmp_path, monkeypatch):
        """目录不可写时降级 _NullStore，不崩。"""
        monkeypatch.setattr(history, "_db_path",
                            lambda: tmp_path / "ro" / "h.db")
        store = history.get_store()
        assert store.add("a", "b", None, "x", "y") == -1
        assert store.recent() == []


class TestXlsxExport:
    def test_export_xlsx_real_datetime_cells(self, tmp_path):
        """xlsx 时间列必须是真日期单元格（Excel 显示可控），CSV 会被
        Excel 按区域设置自作主张转换（用户实测问题）。"""
        from openpyxl import load_workbook
        s = history.HistoryStore(tmp_path / "h.db")
        s.add("zh", "de", None, "测试", "Test", 100)
        s.add("en", "zh", None, "b", "乙", None)
        out = tmp_path / "out.xlsx"
        n = s.export_xlsx(out)
        assert n == 2
        ws = load_workbook(out).active
        rows = list(ws.iter_rows())
        assert rows[0][0].value == "时间"
        from datetime import datetime
        for row in rows[1:]:
            c = row[0]
            assert isinstance(c.value, datetime), "时间列必须是 datetime"
            assert c.number_format == "yyyy-mm-dd hh:mm:ss"
        # 空延迟：写入空串，openpyxl 读回 None（Excel 显示为空，正常）
        assert rows[2][5].value in ("", None)
        s.close()
